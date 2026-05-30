from pathlib import Path

from openpyxl import load_workbook

from company_brain.metadata import make_doc_id
from company_brain.model import Document


def _clean(v) -> str:
    return v.strip() if isinstance(v, str) else ("" if v is None else str(v))


def load_structured(path: str, *, domain: str, sensitivity: str) -> Document:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    records: list[dict] = []
    current_reg = ""
    started = False
    try:
        for row in ws.iter_rows(values_only=True):
            a = _clean(row[0]) if len(row) > 0 else ""
            b = _clean(row[1]) if len(row) > 1 else ""
            if not started:
                if a == "Regulation" and b == "Data attribute":
                    started = True
                continue
            if a:
                current_reg = a          # forward-fill regulation across its attribute rows
            if b and current_reg:
                records.append({"regulation": current_reg, "attribute": b})
    finally:
        wb.close()
    return Document(
        doc_id=make_doc_id(path),
        source_path=path,
        source_type="xlsx",
        title=Path(path).name,
        domain=domain,
        sensitivity=sensitivity,
        records=records,
    )
